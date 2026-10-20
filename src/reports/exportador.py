"""Informes Excel/PDF sin dependencia de la interfaz ni escrituras en PostgreSQL."""
import os
import re
import tempfile
import unicodedata
from pathlib import Path
from xml.sax.saxutils import escape


ENCABEZADOS = ["#", "Apellido", "Nombre", "DNI", "Libreta", "Entrada", "Salida"]
SIN_ASISTENCIAS = "No se registraron asistencias para esta Jornada."


def hora(valor):
    return valor.strftime("%H:%M") if valor is not None else "Sin registrar"


def contador(cantidad):
    if cantidad == 1:
        return "Se registró 1 alumno presente en total."
    return f"Se registraron {cantidad} alumnos presentes en total."


def nombre_sugerido(jornada, formato):
    catedra = unicodedata.normalize("NFKD", jornada.get("catedra") or "sin-catedra")
    catedra = catedra.encode("ascii", "ignore").decode().lower()
    catedra = re.sub(r"[^a-z0-9]+", "-", catedra).strip("-")[:70] or "jornada"
    return f"asistencia_{jornada['fecha']:%Y-%m-%d}_{catedra}.{formato}"


def datos_jornada(jornada):
    return [
        ("Fecha", jornada["fecha"].strftime("%d/%m/%Y")),
        ("Cátedra", jornada.get("catedra") or "Sin cátedra"),
        ("Horario de entrada", hora(jornada["horario_entrada"])),
        ("Horario de salida", hora(jornada["horario_salida"])),
        ("Usuario", jornada["nombre_usuario"]),
    ]


def filas_alumnos(alumnos):
    for indice, alumno in enumerate(alumnos, 1):
        yield [indice, alumno["apellido"], alumno["nombre"], alumno["dni"],
               alumno["libreta"], hora(alumno["entrada"]), hora(alumno["salida"])]


def generar_excel(ruta, jornada, alumnos):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.page import PageMargins

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Asistencia"

    def texto(fila, columna, valor):
        celda = hoja.cell(fila, columna, str(valor))
        # Los nombres y documentos son texto, incluso si comienzan por "=".
        celda.data_type = "s"
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        return celda

    hoja.merge_cells("A1:G1")
    texto(1, 1, "PLANILLA DE ASISTENCIA").font = Font(size=18, bold=True, color="15324F")
    hoja.row_dimensions[1].height = 34
    for fila, (etiqueta, valor) in enumerate(datos_jornada(jornada), 3):
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
        hoja.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=7)
        texto(fila, 1, etiqueta).font = Font(bold=True)
        texto(fila, 3, valor)
        hoja.row_dimensions[fila].height = max(24, 15 * ((len(valor) // 70) + 1))
    for columna, titulo in enumerate(ENCABEZADOS, 1):
        celda = texto(9, columna, titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="15324F")
    hoja.row_dimensions[9].height = 26
    for fila, valores in enumerate(filas_alumnos(alumnos), 10):
        for columna, valor in enumerate(valores, 1):
            celda = texto(fila, columna, valor)
            if fila % 2 == 0:
                celda.fill = PatternFill("solid", fgColor="EEF3F8")
        hoja.row_dimensions[fila].height = max(
            28, 15 * (1 + max(len(str(v)) // 22 for v in valores))
        )
    final = 10 + len(alumnos)
    if not alumnos:
        hoja.merge_cells("A10:G10")
        texto(10, 1, SIN_ASISTENCIAS)
        hoja.row_dimensions[10].height = 28
        final = 11
    hoja.merge_cells(start_row=final+1, start_column=1, end_row=final+1, end_column=7)
    texto(final+1, 1, contador(len(alumnos))).font = Font(bold=True)
    for columna, ancho in zip("ABCDEFG", [6, 26, 26, 18, 18, 16, 18]):
        hoja.column_dimensions[columna].width = ancho
    hoja.freeze_panes = "C10"
    if alumnos:
        hoja.auto_filter.ref = f"A9:G{9+len(alumnos)}"
    hoja.print_title_rows = "1:9"
    hoja.print_area = f"A1:G{final+1}"
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4)
    libro.save(ruta)
    libro.close()


def generar_pdf(ruta, jornada, alumnos):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    estilos = getSampleStyleSheet()
    normal = ParagraphStyle("Celda", parent=estilos["Normal"], fontSize=9, leading=12)
    blanco = ParagraphStyle("Encabezado", parent=normal, textColor=colors.white)
    titulo = ParagraphStyle("Titulo", parent=estilos["Title"], textColor=colors.HexColor("#15324F"))
    centro = ParagraphStyle("Indice", parent=normal, alignment=TA_CENTER)
    def parrafo(valor, estilo=normal):
        return Paragraph(escape(str(valor)), estilo)

    documento = SimpleDocTemplate(str(ruta), pagesize=landscape(A4),
                                  rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=36)
    contenido = [parrafo("PLANILLA DE ASISTENCIA", titulo), Spacer(1, 12)]
    for etiqueta, valor in datos_jornada(jornada):
        contenido.append(Paragraph(f"<b>{escape(etiqueta)}:</b> {escape(valor)}", normal))
        contenido.append(Spacer(1, 4))
    contenido.append(Spacer(1, 14))
    filas = [[parrafo(c, blanco) for c in ENCABEZADOS]]
    for valores in filas_alumnos(alumnos):
        filas.append([parrafo(v, centro if i == 0 else normal) for i, v in enumerate(valores)])
    tabla = Table(filas, colWidths=[30, 160, 160, 110, 110, 80, 80], repeatRows=1, hAlign="LEFT")
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#15324F")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LINEBELOW", (0, 0), (-1, 0), .5, colors.HexColor("#15324F")),
    ]))
    contenido.append(tabla)
    if not alumnos:
        contenido.extend([Spacer(1, 12), parrafo(SIN_ASISTENCIAS)])
    contenido.extend([Spacer(1, 16), parrafo(contador(len(alumnos)), estilos["Heading3"])])
    def pie(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#596C80"))
        canvas.drawRightString(landscape(A4)[0] - 30, 20, f"Página {doc.page}")
        canvas.restoreState()
    documento.build(contenido, onFirstPage=pie, onLaterPages=pie)


def exportar_informe(ruta, formato, jornada, alumnos):
    """Reemplazar el destino sólo después de generar un archivo completo."""
    generadores = {"xlsx": generar_excel, "pdf": generar_pdf}
    generador = generadores[formato]
    ruta = Path(ruta)
    descriptor, temporal = tempfile.mkstemp(prefix=".asistencia-", suffix=f".{formato}", dir=ruta.parent)
    os.close(descriptor)
    try:
        generador(temporal, jornada, alumnos)
        os.replace(temporal, ruta)
    finally:
        if os.path.exists(temporal):
            os.unlink(temporal)
    return str(ruta)
